import os
import random
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference

cwd = os.getcwd()
os.chdir(cwd + "/Data Collection and Storage/Excel Spreadsheets in Python")

workbook = Workbook()
sheet = workbook.active


# Set up sample data
rows = [
    [
        "",
        "January",
        "February",
        "March",
        "April",
        "May",
        "June",
        "July",
        "August",
        "September",
        "October",
        "November",
        "December",
    ],
    [1, ""],
    [2, ""],
    [3, ""],
]

for row in rows:
    sheet.append(row)

for row in sheet.iter_rows(min_row=2, max_row=4, min_col=2, max_col=13):
    for cell in row:
        cell.value = random.randrange(5, 100)


# Create Line Chart
chart = LineChart()
data = Reference(worksheet=sheet, min_row=2, max_row=4, min_col=1, max_col=13)

chart.add_data(data, from_rows=True, titles_from_data=True)

# add chart categories
cats = Reference(worksheet=sheet, min_row=1, max_row=1, min_col=2, max_col=13)
chart.set_categories(cats)

# add axis labels
chart.x_axis.title = "Months"
chart.y_axis.title = "Sales (per unit)"

sheet.add_chart(chart, "C6")

workbook.save(filename="line_chart.xlsx")
